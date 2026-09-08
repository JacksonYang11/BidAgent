from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def export_word(draft, path):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.3)
    section.left_margin = section.right_margin = Cm(2.4)
    for style_name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet'):
        style = doc.styles[style_name]
        style.font.name = 'Arial'
        fonts = style.element.get_or_add_rPr().rFonts
        for attribute in list(fonts.attrib):
            if 'theme' in attribute.lower():
                del fonts.attrib[attribute]
        fonts.set(qn('w:eastAsia'), 'SimSun')
        style.font.color.rgb = RGBColor(0, 0, 0)
        for border in style.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    doc.styles['Title'].font.size = Pt(24)
    doc.styles['Title'].paragraph_format.space_before = Pt(0)
    doc.styles['Title'].paragraph_format.space_after = Pt(14)
    doc.styles['Heading 1'].font.size = Pt(16)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(16)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(8)
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.4
    doc.styles['Normal'].paragraph_format.space_after = Pt(6)
    doc.add_heading(draft['title'], 0)
    doc.add_paragraph('投标方案初稿')
    profile = draft['profile']
    doc.add_paragraph('投标单位：' + (profile['name'] or '【待补充：投标企业名称】'))
    doc.add_paragraph('本文件为辅助起草内容，须核对招标原文并补齐证明材料，经人工审核后使用。')
    if profile['is_demo']:
        doc.add_paragraph('本文件含虚构示例企业资料，不可直接用于真实投标。')
    if draft.get('warnings'):
        doc.add_heading('编制待核对事项', 1)
        for warning in draft['warnings'][:30]:
            doc.add_paragraph(warning)
    for chapter in draft['chapters']:
        doc.add_heading(chapter['title'], 1)
        for block in chapter['blocks']:
            if block['type'] == 'paragraph':
                doc.add_paragraph(block['text'])
            elif block['type'] == 'list':
                for item in block['items']:
                    doc.add_paragraph(item, style='List Bullet')
            elif block['type'] == 'table':
                cols = max(len(block['headers']), max((len(r) for r in block['rows']), default=0))
                if not cols:
                    continue
                table = doc.add_table(rows=1, cols=cols)
                table.style = 'Table Grid'
                for index, value in enumerate(block['headers']):
                    table.rows[0].cells[index].text = value
                trpr = table.rows[0]._tr.get_or_add_trPr()
                repeat = OxmlElement('w:tblHeader')
                trpr.append(repeat)
                for cell in table.rows[0].cells:
                    shade = OxmlElement('w:shd')
                    shade.set(qn('w:fill'), 'E8EEF4')
                    cell._tc.get_or_add_tcPr().append(shade)
                borders = OxmlElement('w:tblBorders')
                for side in ('top','left','bottom','right','insideH','insideV'):
                    border = OxmlElement('w:'+side)
                    for key,value in (('val','single'),('sz','4'),('color','D9D9D9')):
                        border.set(qn('w:'+key),value)
                    borders.append(border)
                table._tbl.tblPr.append(borders)
                for row in block['rows']:
                    cells = table.add_row().cells
                    for index, value in enumerate(row[:cols]):
                        cells[index].text = value
                doc.add_paragraph()
        for pending in chapter.get('pending', []):
            doc.add_paragraph('【待补充】' + pending)
    doc.add_heading('响应覆盖清单', 1)
    for row in draft.get('coverage', []):
        doc.add_paragraph(f'{row["name"]}：{row["status"]}')
    doc.core_properties.author = 'WhaleTalk'
    doc.core_properties.last_modified_by = 'WhaleTalk'
    doc.core_properties.title = draft['title']
    doc.save(path)
    # Reopen the OOXML package to catch corrupt exports before offering downloads.
    Document(path)
